# CARINA (Controlled Artificial Road-traffic Intelligence Network Architecture) is an open-source AI ecosystem for real-time, adaptive control of urban traffic light networks.
# Copyright (C) 2026 Gabriel Moraes - Noxfort Systems
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# File: src/blocks/docx_table_builder.py
# Author: Gabriel Moraes
# Date: July 25, 2026

import re
from typing import Dict, Any, List

try:
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
except ImportError:
    pass

from blocks.math_cleaner import clean_latex_math
from blocks.docx_text_builder import add_formatted_text_to_paragraph

def render_markdown_table(doc, table_lines: list, font_name: str = "Arial", font_size: float = 9.5):
    """Parses markdown table lines and builds a styled graphic docx Table with 0.5pt borders and shaded header."""
    rows_data = []
    for line in table_lines:
        clean_l = line.strip()
        # Ignore separator lines like |---|---|
        if re.match(r'^\|[\s:\|-]+\|?$', clean_l):
            continue
        cells = [c.replace('|', '').strip() for c in clean_l.split('|')]
        if cells and cells[0] == '':
            cells = cells[1:]
        if cells and cells[-1] == '':
            cells = cells[:-1]
        if cells:
            rows_data.append(cells)

    if not rows_data:
        return

    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = 'Table Grid'
    table.autofit = True

    # Apply clean 0.5pt thin table grid borders
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders_xml = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="B0BEC5"/>\n'
            f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="B0BEC5"/>\n'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="B0BEC5"/>\n'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="B0BEC5"/>\n'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>\n'
            f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders_xml)

    for row_idx, row in enumerate(rows_data):
        for col_idx, cell_text in enumerate(row):
            if col_idx < num_cols:
                cell = table.cell(row_idx, col_idx)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)

                clean_cell_text = clean_latex_math(cell_text).replace("|", "").strip()

                # Align numeric or short ID columns to center, text columns to left
                clean_num_check = clean_cell_text.replace('.', '', 1).replace(',', '', 1).strip()
                if clean_num_check.isdigit() or len(clean_cell_text) <= 6:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                # Header row styling (ABNT standard - #F2F2F2 fill)
                if row_idx == 0:
                    set_cell_background(cell, "F2F2F2")  # Light gray header fill
                    tcPr = cell._element.get_or_add_tcPr()
                    header_border = parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>\n'
                        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="222222"/>\n'
                        f'</w:tcBorders>'
                    )
                    tcPr.append(header_border)
                    add_formatted_text_to_paragraph(p, clean_cell_text, font_size=font_size, bold=True, font_name=font_name)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")  # Subtle alternating row tint
                    add_formatted_text_to_paragraph(p, clean_cell_text, font_size=font_size, font_name=font_name)

def set_cell_background(cell, fill_hex: str):
    """Applies a background fill color to a docx table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)
